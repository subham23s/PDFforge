import os
import io
import threading
import webbrowser
import tempfile
from flask import Flask, request, send_file, jsonify, render_template
from werkzeug.utils import secure_filename

app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 50 * 1024 * 1024  # 50MB limit

ALLOWED_EXTENSIONS = {
    'txt', 'md', 'html', 'htm',
    'png', 'jpg', 'jpeg', 'gif', 'bmp', 'webp', 'tiff',
    'docx', 'xlsx', 'csv',
    'pdf'
}

PORT = 5050


# ─── Helpers ──────────────────────────────────────────────────────────────────

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS


# ─── Converters ───────────────────────────────────────────────────────────────

def convert_to_pdf(file, filename):
    ext = filename.rsplit('.', 1)[1].lower()

    # ── PDF passthrough ──
    if ext == 'pdf':
        return file.read()

    # ── Images ──
    elif ext in ('png', 'jpg', 'jpeg', 'gif', 'bmp', 'webp', 'tiff'):
        from PIL import Image
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas

        img = Image.open(file)
        if img.mode in ('RGBA', 'P', 'LA'):
            img = img.convert('RGB')

        with tempfile.NamedTemporaryFile(suffix='.jpg', delete=False) as tmp:
            img.save(tmp.name, 'JPEG', quality=95)
            tmp_path = tmp.name

        pdf_buf = io.BytesIO()
        page_w, page_h = A4
        ratio = min(page_w / img.width, page_h / img.height) * 0.9
        new_w, new_h = img.width * ratio, img.height * ratio
        x = (page_w - new_w) / 2
        y = (page_h - new_h) / 2

        c = canvas.Canvas(pdf_buf, pagesize=A4)
        c.drawImage(tmp_path, x, y, new_w, new_h)
        c.save()
        os.unlink(tmp_path)
        return pdf_buf.getvalue()

    # ── Plain text ──
    elif ext == 'txt':
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

        content = file.read().decode('utf-8', errors='replace')
        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=A4,
                                leftMargin=20*mm, rightMargin=20*mm,
                                topMargin=20*mm, bottomMargin=20*mm)
        styles = getSampleStyleSheet()
        story = []
        for line in content.split('\n'):
            if line.strip():
                safe = line.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                story.append(Paragraph(safe, styles['Normal']))
            else:
                story.append(Spacer(1, 6))
        doc.build(story)
        return buf.getvalue()

    # ── Markdown ──
    elif ext == 'md':
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

        content = file.read().decode('utf-8', errors='replace')
        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=A4,
                                leftMargin=20*mm, rightMargin=20*mm,
                                topMargin=20*mm, bottomMargin=20*mm)
        styles = getSampleStyleSheet()
        story = []
        for line in content.split('\n'):
            line = line.strip()
            if line.startswith('# '):
                story.append(Paragraph(line[2:], styles['h1']))
            elif line.startswith('## '):
                story.append(Paragraph(line[3:], styles['h2']))
            elif line.startswith('### '):
                story.append(Paragraph(line[4:], styles['h3']))
            elif line:
                safe = line.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                story.append(Paragraph(safe, styles['Normal']))
            else:
                story.append(Spacer(1, 6))
        doc.build(story)
        return buf.getvalue()

    # ── HTML ──
    elif ext in ('html', 'htm'):
        import re
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

        content = file.read().decode('utf-8', errors='replace')
        text = re.sub(r'<[^>]+>', ' ', content)
        text = (text.replace('&nbsp;', ' ').replace('&amp;', '&')
                    .replace('&lt;', '<').replace('&gt;', '>').replace('&quot;', '"'))

        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=A4,
                                leftMargin=20*mm, rightMargin=20*mm,
                                topMargin=20*mm, bottomMargin=20*mm)
        styles = getSampleStyleSheet()
        story = []
        for line in text.split('\n'):
            line = line.strip()
            if line:
                safe = line.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                story.append(Paragraph(safe, styles['Normal']))
            else:
                story.append(Spacer(1, 4))
        doc.build(story)
        return buf.getvalue()

    # ── Word DOCX ──
    elif ext == 'docx':
        from docx import Document
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

        with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tmp:
            tmp.write(file.read())
            tmp_path = tmp.name

        doc = Document(tmp_path)
        os.unlink(tmp_path)

        buf = io.BytesIO()
        styles = getSampleStyleSheet()
        pdf_doc = SimpleDocTemplate(buf, pagesize=A4,
                                    leftMargin=20*mm, rightMargin=20*mm,
                                    topMargin=20*mm, bottomMargin=20*mm)
        story = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                safe = text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                if para.style.name.startswith('Heading 1'):
                    story.append(Paragraph(safe, styles['h1']))
                elif para.style.name.startswith('Heading 2'):
                    story.append(Paragraph(safe, styles['h2']))
                elif para.style.name.startswith('Heading 3'):
                    story.append(Paragraph(safe, styles['h3']))
                else:
                    story.append(Paragraph(safe, styles['Normal']))
            else:
                story.append(Spacer(1, 6))
        pdf_doc.build(story)
        return buf.getvalue()

    # ── Excel XLSX ──
    elif ext == 'xlsx':
        import openpyxl
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
            tmp.write(file.read())
            tmp_path = tmp.name

        wb = openpyxl.load_workbook(tmp_path, read_only=True)
        os.unlink(tmp_path)

        buf = io.BytesIO()
        styles = getSampleStyleSheet()
        pdf_doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                                    leftMargin=15*mm, rightMargin=15*mm,
                                    topMargin=15*mm, bottomMargin=15*mm)
        story = []
        for sheet in wb.worksheets:
            story.append(Paragraph(f"Sheet: {sheet.title}", styles['h2']))
            story.append(Spacer(1, 6))
            data = [
                [str(cell) if cell is not None else '' for cell in row]
                for row in sheet.iter_rows(values_only=True)
            ]
            if data:
                t = Table(data)
                t.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2d2d2d')),
                    ('TEXTCOLOR',  (0, 0), (-1, 0), colors.white),
                    ('FONTSIZE',   (0, 0), (-1, -1), 8),
                    ('GRID',       (0, 0), (-1, -1), 0.5, colors.grey),
                    ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f5f5f5')]),
                    ('PADDING',    (0, 0), (-1, -1), 4),
                ]))
                story.append(t)
                story.append(Spacer(1, 12))
        pdf_doc.build(story)
        return buf.getvalue()

    # ── CSV ──
    elif ext == 'csv':
        import csv
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.units import mm
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

        content = file.read().decode('utf-8', errors='replace')
        data = list(csv.reader(io.StringIO(content)))

        buf = io.BytesIO()
        pdf_doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                                    leftMargin=15*mm, rightMargin=15*mm,
                                    topMargin=15*mm, bottomMargin=15*mm)
        story = []
        if data:
            t = Table(data)
            t.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1a1a2e')),
                ('TEXTCOLOR',  (0, 0), (-1, 0), colors.white),
                ('FONTSIZE',   (0, 0), (-1, -1), 8),
                ('GRID',       (0, 0), (-1, -1), 0.5, colors.grey),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f0f0f0')]),
                ('PADDING',    (0, 0), (-1, -1), 4),
            ]))
            story.append(t)
        pdf_doc.build(story)
        return buf.getvalue()

    else:
        raise ValueError(f"Unsupported format: .{ext}")


# ─── Routes ───────────────────────────────────────────────────────────────────

@app.route('/')
def index():
    return render_template('index.html')


@app.route('/convert', methods=['POST'])
def convert():
    if 'file' not in request.files:
        return jsonify({'error': 'No file uploaded'}), 400

    file = request.files['file']
    if not file.filename:
        return jsonify({'error': 'No file selected'}), 400
    if not allowed_file(file.filename):
        return jsonify({'error': 'Unsupported file type'}), 400

    try:
        pdf_bytes = convert_to_pdf(file.stream, secure_filename(file.filename))
        buf = io.BytesIO(pdf_bytes)
        buf.seek(0)
        out_name = file.filename.rsplit('.', 1)[0] + '.pdf'
        return send_file(buf, mimetype='application/pdf',
                         as_attachment=True,
                         download_name=out_name)
    except Exception as e:
        return jsonify({'error': str(e)}), 500


# ─── Auto open browser ────────────────────────────────────────────────────────

def open_browser():
    webbrowser.open(f'http://127.0.0.1:{PORT}')


# ─── Entry point ──────────────────────────────────────────────────────────────

if __name__ == '__main__':
    print(f"\n  ⚡ PDFforge running at http://127.0.0.1:{PORT}\n")
    # Open browser after 1 second (gives Flask time to start)
    threading.Timer(1.0, open_browser).start()
    app.run(debug=False, host='0.0.0.0', port=PORT)